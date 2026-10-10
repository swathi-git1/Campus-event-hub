import docx

doc = docx.Document(r'Campus_Event_Hub_Report\Final_Report.docx')

print(f"Total paragraphs in generated docx: {len(doc.paragraphs)}")

for i, p in enumerate(doc.paragraphs):
    has_img = ''
    if 'graphic' in p._p.xml:
        for node in p._p.xpath('.//wp:extent'):
            w = int(node.get('cx', 0)) / 914400.0
            h = int(node.get('cy', 0)) / 914400.0
            has_img = f'[IMG: {w:.2f} x {h:.2f} in]'
    txt = p.text.strip()
    pb = '[PAGE BREAK]' if 'w:br w:type="page"' in p._p.xml or p.paragraph_format.page_break_before else ''
    if pb or (txt.startswith('CHAPTER') or txt.startswith('8.') or txt.startswith('7.')):
        print(f"P{i:03d} {pb} {has_img}: {txt[:60]}")
