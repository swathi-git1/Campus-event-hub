import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

REPORT_DIR = 'Campus_Event_Hub_Report'
FIG_DIR = os.path.join(REPORT_DIR, 'figures')
SUPPORT_DIR = os.path.join(REPORT_DIR, 'supporting_files')
OUTPUT_DOCX = os.path.join(REPORT_DIR, 'Final_Report.docx')

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=60, bottom=60, left=90, right=90):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="B0C4DE", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_p(doc, text="", font_size=12, bold=False, italic=False, underline=False,
          align=WD_ALIGN_PARAGRAPH.LEFT, space_before=0, space_after=4, line_spacing=1.18,
          page_break_before=False, keep_with_next=False):
    p = doc.add_paragraph()
    p.alignment = align
    if space_before:
        p.paragraph_format.space_before = Pt(space_before)
    if space_after:
        p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    if page_break_before:
        p.paragraph_format.page_break_before = True
    if keep_with_next:
        p.paragraph_format.keep_with_next = True
        
    if text:
        r = p.add_run(text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(font_size)
        r.bold = bold
        r.italic = italic
        r.underline = underline
    return p

def add_bullet_item(doc, lead_bold, text_rest, space_after=2):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    
    r_bullet = p.add_run("• ")
    r_bullet.font.name = "Times New Roman"
    r_bullet.font.size = Pt(12)
    
    if lead_bold:
        r_lead = p.add_run(lead_bold)
        r_lead.font.name = "Times New Roman"
        r_lead.font.size = Pt(12)
        r_lead.bold = True
        
    r_rest = p.add_run(text_rest)
    r_rest.font.name = "Times New Roman"
    r_rest.font.size = Pt(12)
    return p

def add_compact_table(doc, headers, data, col_widths=None, space_after=4):
    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, color="B0C4DE", sz="4")
    
    # Header row
    hdr_row = table.rows[0]
    for idx, heading in enumerate(headers):
        cell = hdr_row.cells[idx]
        set_cell_background(cell, "D9E2F3")
        set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(heading)
        r.font.name = "Times New Roman"
        r.font.size = Pt(10)
        r.bold = True
        
    # Data rows
    for r_idx, row_data in enumerate(data):
        row = table.rows[r_idx + 1]
        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            set_cell_margins(cell, top=50, bottom=50, left=80, right=80)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(val))
            r.font.name = "Times New Roman"
            r.font.size = Pt(9.5)
            
    # Set widths
    if col_widths:
        for idx, width in enumerate(col_widths):
            for row in table.rows:
                row.cells[idx].width = Inches(width)
                
    return table

def add_figure(doc, img_name, caption_text="", expl_text=None, width_in=4.4, space_before=4, space_after=4, keep_with_next=True, page_break_before=False):
    img_path = os.path.join(FIG_DIR, img_name)
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(space_before)
        p_img.paragraph_format.space_after = Pt(2)
        if page_break_before:
            p_img.paragraph_format.page_break_before = True
        if keep_with_next:
            p_img.paragraph_format.keep_with_next = True
        r_img = p_img.add_run()
        r_img.add_picture(img_path, width=Inches(width_in))
        
        if caption_text:
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_before = Pt(2)
            p_cap.paragraph_format.space_after = Pt(3)
            if expl_text or keep_with_next:
                p_cap.paragraph_format.keep_with_next = True
            r_cap = p_cap.add_run(caption_text)
            r_cap.font.name = "Times New Roman"
            r_cap.font.size = Pt(10)
            r_cap.italic = True
            
        if expl_text:
            add_p(doc, expl_text, font_size=11, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=1, space_after=space_after, line_spacing=1.15)
    else:
        print(f"WARNING: Image not found: {img_path}")

print("Helper functions verified.")
