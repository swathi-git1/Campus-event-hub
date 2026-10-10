import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

REPORT_DIR = 'Campus_Event_Hub_Report'
FIG_DIR = os.path.join(REPORT_DIR, 'figures')
SUPPORT_DIR = os.path.join(REPORT_DIR, 'supporting_files')
OUTPUT_DOCX = os.path.join(REPORT_DIR, 'Final_Report.docx')

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="D1D5DB", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_para(doc, text="", style=None, font_size=12, bold=False, italic=False, underline=False,
             align=WD_ALIGN_PARAGRAPH.LEFT, space_before=0, space_after=6, line_spacing=1.15):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before) if space_before else None
    p.paragraph_format.space_after = Pt(space_after) if space_after else None
    p.paragraph_format.line_spacing = line_spacing
    
    if text:
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(font_size)
        run.bold = bold
        run.italic = italic
        run.underline = underline
    return p

def add_heading1(doc, title, chapter_num=None):
    if chapter_num is not None:
        p_ch = doc.add_paragraph()
        p_ch.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_ch.paragraph_format.space_before = Pt(12)
        p_ch.paragraph_format.space_after = Pt(6)
        r_ch = p_ch.add_run(f"CHAPTER {chapter_num}")
        r_ch.font.name = "Times New Roman"
        r_ch.font.size = Pt(16)
        r_ch.bold = True
        
    p_t = doc.add_paragraph()
    p_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_t.paragraph_format.space_before = Pt(6)
    p_t.paragraph_format.space_after = Pt(18)
    r_t = p_t.add_run(title)
    r_t.font.name = "Times New Roman"
    r_t.font.size = Pt(16)
    r_t.bold = True
    return p_t

def add_heading2(doc, title):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(title)
    r.font.name = "Times New Roman"
    r.font.size = Pt(13)
    r.bold = True
    return p

def add_heading3(doc, title):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title)
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.bold = True
    r.underline = True
    return p

def add_bullet(doc, lead_bold, text_rest):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.2
    
    r_bullet = p.add_run("•  ")
    r_bullet.font.name = "Times New Roman"
    r_bullet.font.size = Pt(12)
    
    if lead_bold:
        r_lead = p.add_run(lead_bold)
        r_lead.font.name = "Times New Roman"
        r_lead.font.size = Pt(12)
        r_lead.bold = True
        
    r_rest = p.add_run(text_rest)
    r_rest.font.name = "Times New Roman"
    r_rest.font.size = Pt(12)
    return p

def add_custom_table(doc, headers, data, col_widths=None):
    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, color="B0C4DE", sz="4")
    
    # Header row
    hdr_row = table.rows[0]
    for idx, heading in enumerate(headers):
        cell = hdr_row.cells[idx]
        set_cell_background(cell, "D9E2F3")
        set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(heading)
        r.font.name = "Times New Roman"
        r.font.size = Pt(11)
        r.bold = True
        
    # Data rows
    for r_idx, row_data in enumerate(data):
        row = table.rows[r_idx + 1]
        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(val))
            r.font.name = "Times New Roman"
            r.font.size = Pt(11)
            
    # Set widths
    if col_widths:
        for idx, width in enumerate(col_widths):
            for row in table.rows:
                row.cells[idx].width = Inches(width)
                
    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(0)
    p_spacer.paragraph_format.space_after = Pt(6)
    return table

def add_image_with_caption(doc, img_path, caption_text, width_inches=5.8, is_fig=True, space_after=10):
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(6)
        p_img.paragraph_format.space_after = Pt(2)
        r_img = p_img.add_run()
        r_img.add_picture(img_path, width=Inches(width_inches))
        
        if caption_text:
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_before = Pt(3)
            p_cap.paragraph_format.space_after = Pt(space_after)
            r_cap = p_cap.add_run(caption_text)
            r_cap.font.name = "Times New Roman"
            r_cap.font.size = Pt(10)
            r_cap.italic = True
    else:
        print(f"WARNING: Image not found: {img_path}")

print("Helper functions compiled successfully.")
