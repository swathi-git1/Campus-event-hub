import docx

doc = docx.Document(r'c:\Users\S.Gomathi\Downloads\Store_Management_System_Report_Revised.docx')

print("=== IMAGE SIZES IN REFERENCE DOCX ===")
for i, p in enumerate(doc.paragraphs):
    if 'graphic' in p._p.xml:
        for node in p._p.xpath('.//wp:extent'):
            w = int(node.get('cx', 0)) / 914400.0
            h = int(node.get('cy', 0)) / 914400.0
            # get following text or prev text
            prev_t = doc.paragraphs[i-1].text.strip() if i > 0 else ''
            next_t = doc.paragraphs[i+1].text.strip() if i < len(doc.paragraphs)-1 else ''
            cap = next_t if (next_t.startswith('fig') or next_t.startswith('Fig')) else prev_t
            print(f"P{i:03d}: {w:.2f} in wide x {h:.2f} in high | {cap[:50]}")
