import docx

doc = docx.Document(r'c:\Users\S.Gomathi\Downloads\Store_Management_System_Report_Revised.docx')

print(f"Total paragraphs in reference docx: {len(doc.paragraphs)}")

breaks = []
for i, p in enumerate(doc.paragraphs):
    if 'w:type="page"' in p._p.xml or 'w:br w:type="page"' in p._p.xml:
        breaks.append((i, p.text.strip()[:50]))

print(f"Paragraphs with page breaks ({len(breaks)}):")
for b in breaks:
    print(f"  P{b[0]}: {b[1]}")

# Let's check image dimensions in reference docx
img_dims = []
for i, p in enumerate(doc.paragraphs):
    if 'graphic' in p._p.xml:
        # find extent cx and cy
        for node in p._p.xpath('.//wp:extent'):
            cx = int(node.get('cx', 0))
            cy = int(node.get('cy', 0))
            # convert EMUs to inches: 1 inch = 914400 EMUs
            w_in = cx / 914400.0
            h_in = cy / 914400.0
            img_dims.append((i, w_in, h_in, p.text.strip()[:40]))

print(f"\nTotal images in reference docx: {len(img_dims)}")
print("Sample of image dimensions (width, height in inches):")
for item in img_dims[:20]:
    print(f"  P{item[0]}: {item[1]:.2f} in x {item[2]:.2f} in | text: {item[3]}")
